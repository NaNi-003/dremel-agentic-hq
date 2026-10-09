import argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_docx(md_filepath, output_filepath):
    doc = Document()
    
    # Set default styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Dremel Colors
    dremel_blue = RGBColor(0x00, 0x5b, 0x9f)
    dremel_orange = RGBColor(0xc6, 0x4d, 0x0d)

    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    
    for line in lines:
        raw_line = line.strip()
        
        if raw_line.startswith('```'):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line.strip('\n'))
            p.style.font.name = 'Courier New'
            p.style.font.size = Pt(10)
            p.style.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            p.paragraph_format.left_indent = Inches(0.5)
            continue

        if not raw_line:
            continue

        if raw_line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(raw_line[2:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = dremel_blue
            
        elif raw_line.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(raw_line[3:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = dremel_orange
            
        elif raw_line.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(raw_line[4:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = dremel_blue
            
        elif raw_line.startswith('#### '):
            p = doc.add_paragraph()
            run = p.add_run(raw_line[5:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            
        elif raw_line.startswith('- ') or raw_line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = raw_line[2:].strip()
            # Simple bold formatting parser
            parts = text.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 != 0:
                    run.bold = True
        
        elif raw_line[0].isdigit() and raw_line[1:3] in ['. ', ' ']:
            p = doc.add_paragraph(style='List Number')
            text = raw_line[3:].strip()
            # Simple bold formatting parser
            parts = text.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 != 0:
                    run.bold = True

        else:
            p = doc.add_paragraph()
            # Simple bold formatting parser
            parts = raw_line.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 != 0:
                    run.bold = True

    doc.save(output_filepath)
    print(f"Generated {output_filepath}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Render a markdown file to docx.")
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    create_docx(args.markdown, args.output)
